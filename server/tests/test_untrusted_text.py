"""Текст постов — только данные (НТ-9, НТ-10): формулы в XLSX не исполняются, ссылки строит
сервис, HTML в тексте отдаётся как текст JSON. Посты-ловушки идут через весь конвейер."""

from __future__ import annotations

import io
import re
import zipfile
from datetime import UTC, datetime
from pathlib import Path

from openpyxl import load_workbook

from mediamonitor.adapters.base import RawPost, SourceInfo
from mediamonitor.adapters.fixture import FixtureStore, Recording
from mediamonitor.normalize.urls import PostRef
from mediamonitor.services.export import safe_text
from tests.conftest import make_app, make_settings, new_workspace, open_client, run_search

SOURCE = SourceInfo(
    platform="vk",
    owner_id="-777",
    slug="evil_public",
    name='=HYPERLINK("https://evil.example";"Нажми")',
    url="https://vk.com/evil_public",
    subscribers=10,
    type="community",
    district="Красноярск",
    regional=True,
    in_catalog=True,
)
TRAPS = [
    '=HYPERLINK("https://evil.example";"Красноярск") формула в начале поста.',
    "+7 (391) 000-00-00 — звоните, Красноярск, формула не нужна.",
    "-1+1 формула в Красноярске.",
    "@SUM(A1:A9) формула, Красноярск.",
    "\tТабуляция и формула =1+1 в Красноярске.",
    "<script>alert('xss')</script> Красноярск, формула и ссылка javascript:alert(1).",
]


def evil_store() -> FixtureStore:
    posts = [
        RawPost(
            ref=PostRef("vk", "-777", str(100 + i)),
            source=SOURCE,
            text=text,
            title=None,
            published_at=datetime(2026, 10, 10, 3, i, tzinfo=UTC),
            time_reliability="exact",
            edited=False,
            repost_of=None,
            has_media=False,
            views=None,
            reactions=None,
            alt_urls=(
                "javascript:alert(1)",
                f"https://m.vk.com/wall-777_{100 + i}",
                "https://evil.example/x",
            ),
        )
        for i, text in enumerate(TRAPS)
    ]
    return FixtureStore(
        sources={"evil": SOURCE},
        recordings=[
            Recording(
                id="evil",
                regex=[re.compile("формул")],
                shingle_index=set(),
                min_shingle_hits=0,
                meaning_id=None,
                posts=posts,
            )
        ],
        posts={p.ref.id: p for p in posts},
        everywhere_pool={"vk": 2, "ok": 1, "max": 1},
    )


async def test_traps_are_plain_text_everywhere(tmp_path: Path) -> None:
    harness = make_app(make_settings(tmp_path), store=evil_store())
    async with open_client(harness.app) as client:
        headers = await new_workspace(client)
        search_id, events = await run_search(client, headers, "формула")
        selection = events[-1]["data"]["selection"]
        assert selection["counts"]["total"] == len(TRAPS)
        for item in selection["items"]:
            publication = item["publication"]
            # Ссылка — только построенная сервисом из ключа поста; чужие адреса отброшены.
            assert re.fullmatch(r"https://vk\.com/wall-777_\d+", publication["url"])
            assert publication["mergedUrls"] == [
                publication["url"].replace("https://vk.com", "https://m.vk.com")
            ]
        texts = [i["publication"]["text"] for i in selection["items"]]
        assert any("<script>" in t for t in texts)  # отдаём как есть, в JSON-строке

        response = await client.get(f"/api/searches/{search_id}/export.xlsx", headers=headers)
        content = response.content
        # В файле нет ни одной формулы.
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            for name in archive.namelist():
                if name.startswith("xl/worksheets/"):
                    assert "<f>" not in archive.read(name).decode("utf-8"), name
        workbook = load_workbook(io.BytesIO(content))
        sheet = workbook["Публикации"]
        dangerous = 0
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                if isinstance(cell.value, str) and cell.value[:1] in "=+-@\t":
                    dangerous += 1
                    assert cell.data_type == "s"
                    assert cell.quotePrefix is True
            source_cell, url_cell = row[2], row[6]
            assert source_cell.value == SOURCE.name  # текст сохранён дословно, но как текст
            assert url_cell.hyperlink is not None
            target = url_cell.hyperlink.target
            assert target is not None
            assert target.startswith("https://vk.com/wall-777_")
        assert dangerous >= len(TRAPS)

        preview = (
            await client.get(f"/api/searches/{search_id}/export/preview", headers=headers)
        ).json()
        for row in preview["publications"]["rows"]:
            assert row[2] == "'" + SOURCE.name  # в предпросмотре — с апострофом, как в мокапе


def test_safe_text() -> None:
    assert safe_text("=1+1") == "'=1+1"
    assert safe_text("@cmd") == "'@cmd"
    assert safe_text("Красноярск") == "Красноярск"
