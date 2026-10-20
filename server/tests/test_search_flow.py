"""Сквозной сценарий на записанных примерах: поиск → поток событий → подборка →
предпросмотр выгрузки → XLSX (Сц. ТЗ 1 и 4), а также повтор, остановка и «Недавние»."""

from __future__ import annotations

import asyncio
import io
import re
from typing import cast
from urllib.parse import quote

import httpx
from openpyxl import load_workbook

from mediamonitor.contract import PlatformId
from mediamonitor.settings import Settings
from tests.conftest import (
    BDD,
    PERIOD,
    Harness,
    make_app,
    new_workspace,
    open_client,
    read_stream,
    run_search,
)


async def test_health(client: httpx.AsyncClient) -> None:
    response = await client.get("/health")
    assert response.status_code == 200
    assert response.json() == {"status": "ok", "database": "ok", "version": "0.1.0"}


async def test_search_to_xlsx(client: httpx.AsyncClient) -> None:
    headers = await new_workspace(client)
    search_id, events = await run_search(client, headers)
    types = [e["event"] for e in events]
    assert types[0] == "started"
    assert types[1] == "ai-unavailable"  # адрес нейросети не задан
    assert types[2] == "topic"
    assert types[-1] == "finished"
    assert "items" in types
    assert "platform" in types
    assert [e["id"] for e in events] == sorted(e["id"] for e in events)

    finished = events[-1]["data"]["selection"]
    assert finished["phase"] == "done"
    assert finished["counts"]["total"] > 10
    assert {s["platform"]: s["state"] for s in finished["statuses"]} == {
        "vk": "found",
        "ok": "found",
        "max": "found",
    }
    # Смысловая проверка недоступна → спорные помечены «не проверено ИИ» (SC-08).
    assert finished["aiAvailable"] is False
    assert any(n["code"] == "ai-unavailable" for n in finished["notices"])
    methods = {i["method"] for i in finished["items"]}
    assert "ai-confirmed" not in methods
    # Числа совпадают: площадки = подборка (правило «числа совпадают»).
    assert sum(s["found"] for s in finished["statuses"]) == finished["counts"]["total"]

    selection = (await client.get(f"/api/searches/{search_id}", headers=headers)).json()
    assert selection["counts"] == finished["counts"]
    for item in selection["items"]:
        publication = item["publication"]
        assert item["regionReason"]["regionalSource"] or item["regionReason"]["places"]
        if publication["text"] is not None and item["excerpt"]:
            excerpt = item["excerpt"]
            assert publication["text"][excerpt["start"] : excerpt["end"]] == excerpt["text"]
            assert 0 <= item["evidence"]["sentenceIndex"] < len(publication["sentences"])
        assert publication["url"].startswith(
            ("https://vk.com/wall", "https://ok.ru/group/", "https://max.ru/")
        )
        assert publication["publishedAt"] is None or publication["publishedAt"].endswith("+07:00")

    preview = (
        await client.get(f"/api/searches/{search_id}/export/preview", headers=headers)
    ).json()
    assert preview["rowCount"] == selection["counts"]["total"]
    assert preview["publications"]["totalRows"] == preview["rowCount"]
    assert preview["timeZoneNote"] == "время красноярское, UTC+7"
    # В имени файла — день и время поиска (по Красноярску), без «:»: две выгрузки за день
    # не путаются.
    assert re.fullmatch(r"Медиарадар_\w+_\d{4}-\d{2}-\d{2}_\d{2}-\d{2}\.xlsx", preview["fileName"])

    response = await client.get(f"/api/searches/{search_id}/export.xlsx", headers=headers)
    assert response.status_code == 200
    assert "attachment" in response.headers["content-disposition"]
    assert quote(preview["fileName"]) in response.headers["content-disposition"]
    # ExportFile.rowCount интерфейса — заголовком; CORS открывает его браузеру.
    assert response.headers["x-export-rows"] == str(preview["rowCount"])
    workbook = load_workbook(io.BytesIO(response.content))
    assert workbook.sheetnames == ["Публикации", "Сводка"]
    sheet = workbook["Публикации"]
    assert sheet.max_row - 1 == preview["rowCount"]
    header = [c.value for c in sheet[1]]
    assert header == [c["title"] for c in preview["publications"]["columns"]]
    # Записанные примеры — демо-публикации: вместо вымышленной ссылки пометка, без гиперссылки
    # (живые ссылки и их гиперссылки — test_untrusted_text.py, test_vk_live.py).
    url_cell = sheet.cell(row=2, column=7)
    assert url_cell.value == "демо-публикация, ссылки нет"
    assert url_cell.hyperlink is None


async def test_release_spread(client: httpx.AsyncClient) -> None:
    from mediamonitor.adapters.fixture import FixtureStore

    store = FixtureStore.load()
    release = next(r for r in store.recordings if r.id == "release")
    original = next(p for p in release.posts if p.ref.owner_id == "-100104")
    assert original.text is not None
    headers = await new_workspace(client)
    search_id, events = await run_search(client, headers, original.text)
    finished = events[-1]["data"]["selection"]
    assert finished["conditions"]["kind"] == "news-text"
    groups = [g for g in finished["groups"] if g["kind"] == "release"]
    assert len(groups) == 1
    spread = (await client.get(f"/api/searches/{search_id}/spread", headers=headers)).json()
    group = groups[0]
    assert spread["totals"]["placements"] == len(group["itemIds"])
    assert len(spread["placements"]) == spread["totals"]["placements"]
    assert spread["original"]["itemId"] == group["originalItemId"]
    assert spread["release"]["fragments"]
    # Сверх контракта ветки (условия, статусы, вся подборка) сервер в JSON не отдаёт.
    assert not {"conditions", "statuses", "selection"} & set(spread)
    assert not {"sourceId", "anomalies", "fragmentIds"} & set(spread["placements"][0])


async def test_attach_snapshot(client: httpx.AsyncClient) -> None:
    headers = await new_workspace(client)
    search_id, _ = await run_search(client, headers)
    events = await read_stream(client, headers, search_id, snapshot="true")
    assert [e["event"] for e in events] == ["started", "topic", "items", "finished"]
    assert len({e["id"] for e in events}) == 1


async def test_reconnect_with_last_event_id(client: httpx.AsyncClient) -> None:
    headers = await new_workspace(client)
    search_id, events = await run_search(client, headers)
    middle = events[len(events) // 2]["id"]
    rest = await read_stream(client, headers, search_id, after=middle)
    assert [e["id"] for e in rest] == [e["id"] for e in events if e["id"] > middle]
    response = await client.get(
        f"/api/searches/{search_id}/events", headers={**headers, "Last-Event-ID": str(middle)}
    )
    assert len(response.text) > 0


async def test_platform_failure_and_retry(settings: Settings) -> None:
    harness = make_app(settings, failures={"max": "unavailable"})
    async with open_client(harness.app) as client:
        headers = await new_workspace(client)
        search_id, events = await run_search(client, headers)
        finished = events[-1]["data"]["selection"]
        statuses = {s["platform"]: s for s in finished["statuses"]}
        assert statuses["max"]["state"] == "error"
        assert statuses["max"]["problem"]["code"] == "unavailable"
        assert statuses["max"]["problem"]["retryable"] is True
        assert statuses["vk"]["state"] == "found"  # сбой одной площадки не роняет остальные
        before = finished["counts"]["total"]

        response = await client.post(
            f"/api/searches/{search_id}/platforms/max/retry", headers=headers
        )
        assert response.status_code == 202, response.text
        ticket = response.json()
        retry_events = await read_stream(client, headers, search_id, after=ticket["after"])
        assert retry_events[-1]["event"] == "finished"
        selection = retry_events[-1]["data"]["selection"]
        statuses = {s["platform"]: s for s in selection["statuses"]}
        assert statuses["max"]["state"] == "found"
        assert selection["counts"]["total"] >= before
        assert sum(s["found"] for s in selection["statuses"]) == selection["counts"]["total"]


async def test_no_access_is_not_retryable(settings: Settings) -> None:
    harness = make_app(settings, failures={"ok": "quota"})
    async with open_client(harness.app) as client:
        headers = await new_workspace(client)
        _, events = await run_search(client, headers)
        statuses = {s["platform"]: s for s in events[-1]["data"]["selection"]["statuses"]}
        assert statuses["ok"]["state"] == "no-access"
        assert statuses["ok"]["problem"] == {
            "code": "quota-exhausted",
            "message": statuses["ok"]["problem"]["message"],
            "retryable": False,
        }


async def test_bad_response_is_error_not_nothing(settings: Settings) -> None:
    harness = make_app(settings, failures={"vk": "bad-response"})
    async with open_client(harness.app) as client:
        headers = await new_workspace(client)
        _, events = await run_search(client, headers)
        statuses = {s["platform"]: s for s in events[-1]["data"]["selection"]["statuses"]}
        assert statuses["vk"]["state"] == "error"  # «тихий ноль» — ошибка (НТ-5)
        assert statuses["vk"]["problem"]["code"] == "bad-response"


async def test_platform_timeout(settings: Settings) -> None:
    fast = settings.model_copy(update={"platform_timeout_sec": 0.3})
    harness = make_app(fast, failures={"ok": "timeout"})
    async with open_client(harness.app) as client:
        headers = await new_workspace(client)
        _, events = await run_search(client, headers)
        statuses = {s["platform"]: s for s in events[-1]["data"]["selection"]["statuses"]}
        assert statuses["ok"]["state"] == "error"
        assert statuses["ok"]["problem"]["code"] == "timeout"


async def test_stop_search(settings: Settings) -> None:
    harness = make_app(settings, step_delay_sec=0.05)
    async with open_client(harness.app) as client:
        headers = await new_workspace(client)
        body = {"text": BDD, "period": PERIOD, "mode": "everywhere"}
        search_id = (await client.post("/api/searches", json=body, headers=headers)).json()[
            "searchId"
        ]
        await asyncio.sleep(0.4)
        response = await client.post(f"/api/searches/{search_id}/stop", headers=headers)
        assert response.status_code == 204
        events = await read_stream(client, headers, search_id)
        assert events[-1]["event"] == "stopped"
        selection = events[-1]["data"]["selection"]
        assert selection["phase"] == "stopped"
        assert selection["conditions"]["finishedAt"] is not None
        assert all(s["state"] in ("stopped", "found", "nothing") for s in selection["statuses"])


async def test_recent_and_repeat(client: httpx.AsyncClient) -> None:
    headers = await new_workspace(client)
    first_id, _ = await run_search(client, headers)
    recent = (await client.get("/api/recent", headers=headers)).json()
    assert [r["searchId"] for r in recent] == [first_id]
    assert recent[0]["request"]["text"] == BDD
    assert recent[0]["total"] > 0
    second_id, events = await run_search(client, headers, repeat_of=first_id)
    assert events[-1]["data"]["selection"]["counts"]["new"] == 0
    recent = (await client.get("/api/recent", headers=headers)).json()
    assert [r["searchId"] for r in recent] == [second_id]  # повтор заменяет запись
    workspace = (await client.get("/api/workspace", headers=headers)).json()
    assert workspace["recentCount"] == 1


async def test_exclusions_and_export_rules(client: httpx.AsyncClient) -> None:
    headers = await new_workspace(client)
    search_id, events = await run_search(client, headers)
    items = events[-1]["data"]["selection"]["items"]
    total = events[-1]["data"]["selection"]["counts"]["total"]
    first = items[0]["id"]
    counts = (
        await client.post(
            f"/api/searches/{search_id}/exclusions",
            json={"itemIds": [first], "excluded": True},
            headers=headers,
        )
    ).json()
    assert counts["total"] == total - 1
    assert counts["excluded"] == 1
    source_id = items[1]["publication"]["source"]["id"]
    result = (
        await client.post(
            f"/api/searches/{search_id}/exclusions/source",
            json={"sourceId": source_id},
            headers=headers,
        )
    ).json()
    assert result["itemIds"]
    assert result["counts"]["excluded"] == 1 + len(result["itemIds"])
    preview = (
        await client.get(f"/api/searches/{search_id}/export/preview", headers=headers)
    ).json()
    assert preview["rowCount"] == result["counts"]["total"]
    assert preview["excludedCount"] == result["counts"]["excluded"]
    filtered = (
        await client.get(
            f"/api/searches/{search_id}/export/preview",
            params={"onlyRelevance": ["direct"]},
            headers=headers,
        )
    ).json()
    assert filtered["appliedFilter"] == ["direct"]
    assert filtered["rowCount"] == result["counts"]["byRelevance"]["direct"]
    undated = (
        await client.get(
            f"/api/searches/{search_id}/export/preview",
            params={"onlyUndated": "true"},
            headers=headers,
        )
    ).json()
    assert "onlyUndated" not in undated
    selection = (await client.get(f"/api/searches/{search_id}", headers=headers)).json()
    undated_items = [i for i in selection["items"] if i["publication"]["publishedAt"] is None]
    assert undated["rowCount"] == sum(1 for i in undated_items if not i["excluded"])
    for row in undated["publications"]["rows"]:
        assert row[4] is None
        assert row[5] == "дата неизвестна: в период не засчитана"
    # Вернуть исключённое.
    counts = (
        await client.post(
            f"/api/searches/{search_id}/exclusions",
            json={"itemIds": [first], "excluded": False},
            headers=headers,
        )
    ).json()
    assert counts["excluded"] == len(result["itemIds"])


async def test_add_by_url_and_recheck_links(harness: Harness, client: httpx.AsyncClient) -> None:
    headers = await new_workspace(client)
    search_id, events = await run_search(client, headers)
    response = await client.post(
        f"/api/searches/{search_id}/items",
        json={"url": "https://m.vk.com/wall-100106_57990"},
        headers=headers,
    )
    assert response.status_code == 201, response.text
    item = response.json()
    assert item["method"] == "manual"
    assert item["id"] == "vk:-100106_57990"
    assert item["publication"]["url"] == "https://vk.com/wall-100106_57990"
    again = await client.post(
        f"/api/searches/{search_id}/items",
        json={"url": "https://vk.com/wall-100106_57990"},
        headers=headers,
    )
    assert again.status_code == 422
    assert again.json()["code"] == "invalid"
    missing = await client.post(
        f"/api/searches/{search_id}/items", json={"url": "https://vk.com/wall-1_1"}, headers=headers
    )
    assert missing.status_code == 404
    assert missing.json()["code"] == "not-found"

    victim = events[-1]["data"]["selection"]["items"][0]["id"]
    platform = cast(PlatformId, victim.split(":")[0])
    harness.adapters[platform].deleted.add(victim)
    report = (await client.post(f"/api/searches/{search_id}/links/recheck", headers=headers)).json()
    assert victim in report["deletedItemIds"]
    assert report["ok"] == report["checked"] - len(report["deletedItemIds"])
    selection = (await client.get(f"/api/searches/{search_id}", headers=headers)).json()
    states = {i["id"]: i["publication"]["linkCheck"]["state"] for i in selection["items"]}
    assert states[victim] == "deleted"
