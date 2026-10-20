"""Воронка «Э-Отбор» на сервере (бэклог п. 45): лёгкая модель смысла вместо BERTA.

Проверяем устройство, а не качество отбора: предложения и доказательства воронки
совпадают с публикацией, числа совпадают везде, векторы кэшируются в базе и не
считаются второй раз, результаты идут в поток порциями, сервис эмбеддингов на
видеокарте и откат на CPU, бюджет нейросети — на весь поиск.
"""

from __future__ import annotations

import io
import json
from pathlib import Path
from typing import Any

import httpx
import numpy as np
import pytest
from mm_matching.text import split_sentences as mm_split
from mm_matching.topic import clear_card_cache
from openpyxl import load_workbook
from sqlalchemy import func, select

from mediamonitor.db.models import PostEmbeddingRow
from mediamonitor.matcher.embeddings import (
    CachingEmbedder,
    EmbeddingServiceError,
    EmbeddingsUnavailableError,
    HttpEmbedder,
    build_embedder,
)
from mediamonitor.matcher.funnel import FunnelMatcher
from mediamonitor.settings import Settings
from tests.conftest import (
    Harness,
    make_app,
    make_settings,
    new_workspace,
    open_client,
    read_stream,
    run_search,
    unique,
)
from tests.contract.test_contract_schema import assert_contract
from tests.fakes import AvailableJudge, FakeEmbedder, FakeLlm, make_funnel

RELEASE_TEXT: str | None = None


def release_text() -> str:
    global RELEASE_TEXT  # noqa: PLW0603
    if RELEASE_TEXT is None:
        data = json.loads(
            (
                Path(__file__).resolve().parents[1] / "src/mediamonitor/fixtures/recordings.json"
            ).read_text("utf-8")
        )
        release = next(r for r in data["recordings"] if r["id"] == "release")
        RELEASE_TEXT = next(p["text"] for p in release["posts"] if p["ownerId"] == "-100104")
    return RELEASE_TEXT


def funnel_app(
    tmp_path: Path, *, embedder: FakeEmbedder | None = None, **settings: Any
) -> tuple[Harness, list[FunnelMatcher]]:
    """Приложение с воронкой. Возвращает и созданную воронку — для проверок счётчиков."""
    built: list[FunnelMatcher] = []
    llm = settings.pop("llm", None)

    def factory(s: Settings) -> FunnelMatcher:
        matcher = make_funnel(s, embedder=embedder, llm=llm)
        built.append(matcher)
        return matcher

    config = make_settings(tmp_path, **settings)
    judge = AvailableJudge() if llm is not None else None
    return make_app(config, matcher_factory=factory, judge=judge), built


def check_item(item: dict[str, Any]) -> None:
    """Доказательство и выдержка — на настоящих предложениях текста (К14)."""
    publication = item["publication"]
    text = publication["text"]
    if text is None:
        return
    expected = [{"index": s.index, "start": s.start, "end": s.end} for s in mm_split(text)]
    assert publication["sentences"] == expected  # то же разбиение, что у воронки
    evidence = item["evidence"]
    assert 0 <= evidence["sentenceIndex"] < len(expected)
    for r in evidence["ranges"]:
        assert 0 <= r["start"] < r["end"] <= len(text)
    excerpt = item["excerpt"]
    assert text[excerpt["start"] : excerpt["end"]] == excerpt["text"]


async def test_topic_search_through_funnel(tmp_path: Path) -> None:
    harness, built = funnel_app(tmp_path)
    async with open_client(harness.app) as client:
        headers = await new_workspace(client)
        search_id, events = await run_search(client, headers)
        types = [e["event"] for e in events]
        assert types[-1] == "finished"
        # Результаты идут в поток порциями по мере отбора (НТ-2), а не одним куском в конце.
        assert types.count("items") >= 3
        for event in events:
            assert_contract("SearchEvent", event["data"])
        topic = next(e["data"]["topic"] for e in events if e["event"] == "topic")
        assert topic["madeBy"] == "rules"  # нейросети нет — карточка воронки по правилам
        selection = (await client.get(f"/api/searches/{search_id}", headers=headers)).json()
        assert_contract("Selection", selection)
        counts = selection["counts"]
        assert counts["total"] > 0
        assert sum(s["found"] for s in selection["statuses"]) == counts["total"]
        for item in selection["items"]:
            check_item(item)
            if item["method"] == "ai-unchecked":
                assert item["relevance"] == "disputed"
                assert "Смысловая проверка недоступна" in item["methodNote"]
        for group in selection["groups"]:
            members = [i for i in selection["items"] if i["groupId"] == group["id"]]
            assert len(members) >= 2
            assert group["kind"] == "copies"
        preview = (
            await client.get(f"/api/searches/{search_id}/export/preview", headers=headers)
        ).json()
        assert preview["rowCount"] == counts["total"]
        xlsx = await client.get(f"/api/searches/{search_id}/export.xlsx", headers=headers)
        assert load_workbook(io.BytesIO(xlsx.content))["Публикации"].max_row - 1 == counts["total"]
        stand = (await client.get("/api/stand", headers=headers)).json()
        assert "Модель смысла fake-berta: CPU сервера" in stand["events"][0]["message"]
    assert built[0].embedder.computed > 0


async def test_release_search_through_funnel(tmp_path: Path) -> None:
    harness, _ = funnel_app(tmp_path)
    async with open_client(harness.app) as client:
        headers = await new_workspace(client)
        search_id, events = await run_search(client, headers, release_text())
        selection = events[-1]["data"]["selection"]
        assert selection["conditions"]["kind"] == "news-text"
        release_groups = [g for g in selection["groups"] if g["kind"] == "release"]
        assert len(release_groups) == 1
        group = release_groups[0]
        spread = (await client.get(f"/api/searches/{search_id}/spread", headers=headers)).json()
        assert_contract("SpreadReport", spread)
        assert spread["totals"]["placements"] == len(group["itemIds"])
        assert spread["original"]["itemId"] == group["originalItemId"]
        fragment_ids = {f["id"] for f in spread["release"]["fragments"]}
        assert fragment_ids
        assert all(f.startswith("p") for f in fragment_ids)
        for item in selection["items"]:
            check_item(item)
            if item["citation"]:
                assert set(item["citation"]["fragmentIds"]) <= fragment_ids
        copies = [i for i in selection["items"] if i["method"] == "copy"]
        assert copies, "перепечатки находятся совпадением фраз"
        assert all(i["placementType"] in ("verbatim", "shortened", "repost") for i in copies)


async def test_vectors_are_cached_in_database(tmp_path: Path) -> None:
    """Второй поиск по тем же постам (новый процесс, та же база) векторы не считает (Р7)."""
    # Своя версия векторов: база может быть общей для прогонов (TEST_DATABASE_URL).
    version = unique()
    first_model = FakeEmbedder()
    harness, built = funnel_app(tmp_path, embedder=first_model, embedding_version=version)
    async with open_client(harness.app) as client:
        headers = await new_workspace(client)
        _, events = await run_search(client, headers)
        total = events[-1]["data"]["selection"]["counts"]["total"]
        ctx = harness.app.state.ctx
        async with ctx.sessions() as session:
            stored = await session.scalar(
                select(func.count())
                .select_from(PostEmbeddingRow)
                .where(PostEmbeddingRow.model == f"fake-berta#{version}")
            )
    first = built[0].embedder.computed
    assert stored
    assert stored >= total

    second_model = FakeEmbedder()
    harness2, built2 = funnel_app(tmp_path, embedder=second_model, embedding_version=version)
    async with open_client(harness2.app) as client:
        headers = await new_workspace(client)
        _, events = await run_search(client, headers)
        assert events[-1]["data"]["selection"]["counts"]["total"] == total
    second = built2[0].embedder.computed
    # Посты из кэша; посчитаны только тексты запроса и граней карточки темы.
    assert second < first
    assert second <= 12
    assert sum(n for _, n in second_model.calls) == second


async def test_new_model_version_recomputes(tmp_path: Path) -> None:
    harness, _ = funnel_app(tmp_path, embedding_version=unique())
    async with open_client(harness.app) as client:
        await run_search(client, await new_workspace(client))
    harness2, built = funnel_app(tmp_path, embedding_version=unique())
    async with open_client(harness2.app) as client:
        await run_search(client, await new_workspace(client))
    assert built[0].embedder.computed > 20  # другая версия — векторы считаются заново


async def test_llm_budget_is_per_search(tmp_path: Path) -> None:
    llm = FakeLlm()
    harness, _ = funnel_app(tmp_path, llm=llm, ollama_url="http://fake", llm_max_gray=3)
    async with open_client(harness.app) as client:
        headers = await new_workspace(client)
        _, events = await run_search(client, headers)
        selection = events[-1]["data"]["selection"]
        confirmed = [i for i in selection["items"] if i["method"] == "ai-confirmed"]
        budget = [
            i
            for i in selection["items"]
            if i["method"] == "ai-unchecked" and "Не хватило времени" in (i["methodNote"] or "")
        ]
        assert 0 < len(confirmed) <= 3
        assert budget, "сверх бюджета спорные остаются «не проверено ИИ»"
        for item in confirmed:
            assert item["evidence"]["basis"] == "ai"
            check_item(item)
    verdict_calls = llm.calls
    assert verdict_calls <= 3 + 1  # три спорных + карточка темы


async def test_llm_down_marks_unchecked(tmp_path: Path) -> None:
    llm = FakeLlm(fail=True)
    harness, _ = funnel_app(tmp_path, llm=llm, ollama_url="http://fake")
    async with open_client(harness.app) as client:
        headers = await new_workspace(client)
        _, events = await run_search(client, headers)
        assert "ai-unavailable" in [e["event"] for e in events]
        selection = events[-1]["data"]["selection"]
        assert selection["aiAvailable"] is False
        unchecked = [i for i in selection["items"] if i["method"] == "ai-unchecked"]
        assert unchecked
        assert all(i["relevance"] == "disputed" for i in unchecked)


async def test_stop_while_matching(tmp_path: Path) -> None:
    harness, _ = funnel_app(tmp_path, match_batch_size=2)
    async with open_client(harness.app) as client:
        headers = await new_workspace(client)
        body = {
            "text": "нагрузочный набор 500",
            "period": {"preset": "custom", "from": "2026-09-20", "to": "2026-10-19"},
            "mode": "everywhere",
        }
        search_id = (await client.post("/api/searches", json=body, headers=headers)).json()[
            "searchId"
        ]
        await client.post(f"/api/searches/{search_id}/stop", headers=headers)
        events = await read_stream(client, headers, search_id)
        assert events[-1]["event"] == "stopped"
        selection = events[-1]["data"]["selection"]
        assert sum(s["found"] for s in selection["statuses"]) == selection["counts"]["total"]


# ---------------------------------------------------------------- сервис эмбеддингов (gpu-node)


def gpu_node(
    *, model: str = "fake-berta", fail: bool = False, seen: list[dict[str, Any]] | None = None
) -> httpx.MockTransport:
    fake = FakeEmbedder()

    def handler(request: httpx.Request) -> httpx.Response:
        if fail:
            raise httpx.ConnectError("нет связи")
        assert request.headers["Authorization"] == "Bearer test-token"
        if request.url.path == "/health":
            return httpx.Response(200, json={"status": "ok", "model": model, "device": "cuda"})
        body = json.loads(request.content)
        if seen is not None:
            seen.append(body)
        assert set(body) == {"texts", "model", "prefix", "normalize"}
        assert body["normalize"] is True
        assert len(body["texts"]) <= 256
        vectors = fake.encode([body["prefix"] + t for t in body["texts"]], "topic")
        return httpx.Response(200, json={"model": model, "dim": 256, "vectors": vectors.tolist()})

    return httpx.MockTransport(handler)


def test_http_embedder_contract() -> None:
    seen: list[dict[str, Any]] = []
    remote = HttpEmbedder(
        "http://gpu.test/", "fake-berta", token="test-token", transport=gpu_node(seen=seen)
    )
    assert remote.health()["device"] == "cuda"
    texts = [f"Новость номер {i} о дорогах Красноярска." for i in range(300)]
    matrix = remote.encode(texts, "topic")
    assert matrix.shape == (300, 256)
    assert np.allclose(np.linalg.norm(matrix, axis=1), 1.0, atol=1e-5)
    assert [len(b["texts"]) for b in seen] == [256, 44]  # пакеты до 256 текстов
    assert {b["prefix"] for b in seen} == {"categorize_topic: "}  # префикс задачи BERTA
    paraphrase = remote.encode(["Текст новости."], "paraphrase")
    assert paraphrase.shape == (1, 256)
    assert seen[-1]["prefix"] == "paraphrase: "


def test_http_embedder_rejects_foreign_model() -> None:
    remote = HttpEmbedder(
        "http://gpu.test", "fake-berta", token="test-token", transport=gpu_node(model="other")
    )
    with pytest.raises(EmbeddingServiceError):
        remote.encode(["текст"], "topic")


def test_fallback_to_cpu_and_back() -> None:
    now = [0.0]
    local = FakeEmbedder()
    broken = HttpEmbedder(
        "http://gpu.test", "fake-berta", token="test-token", transport=gpu_node(fail=True)
    )
    embedder = CachingEmbedder(
        "fake-berta", local=lambda: local, remote=broken, retry_after=60, clock=lambda: now[0]
    )
    embedder.encode(["Первый текст про ДТП."], "topic")
    assert embedder.provider == "fallback"
    assert local.calls, "посчитано на CPU"
    note = embedder.notes()[0]
    assert note.level == "warning"
    assert "недоступен" in note.message
    # Через минуту — снова пробуем сервис; он ожил.
    healthy = HttpEmbedder(
        "http://gpu.test", "fake-berta", token="test-token", transport=gpu_node()
    )
    embedder._remote = healthy
    now[0] = 61.0
    embedder.encode(["Второй текст про гололёд."], "topic")
    assert embedder.provider == "remote"
    assert len(local.calls) == 1
    assert embedder.notes()[0].level == "info"


async def test_search_uses_gpu_node(tmp_path: Path) -> None:
    seen: list[dict[str, Any]] = []
    remote = HttpEmbedder(
        "http://gpu.test", "fake-berta", token="test-token", transport=gpu_node(seen=seen)
    )
    local = FakeEmbedder()

    def factory(settings: Settings) -> FunnelMatcher:
        return make_funnel(settings, embedder=local, remote=remote)

    harness = make_app(
        make_settings(tmp_path, embeddings_url="http://gpu.test"), matcher_factory=factory
    )
    async with open_client(harness.app) as client:
        headers = await new_workspace(client)
        _, events = await run_search(client, headers)
        assert events[-1]["data"]["selection"]["counts"]["total"] > 0
        stand = (await client.get("/api/stand")).json()
    assert seen, "векторы считал сервис на видеокарте"
    assert not local.calls, "CPU сервера не понадобился"
    assert "сервис эмбеддингов" in stand["events"][0]["message"]
    assert stand["events"][0]["level"] == "info"


# ---------------------------------------------------------------- без отката на CPU (слабый сервер)


def token_rejected() -> httpx.MockTransport:
    """Как настоящий gpu-node без нужного токена: /health открыт, /embed — 401."""

    def handler(request: httpx.Request) -> httpx.Response:
        if request.url.path == "/health":
            return httpx.Response(
                200, json={"status": "ok", "model": "fake-berta", "device": "cuda"}
            )
        return httpx.Response(401, json={"detail": "unauthorized"})

    return httpx.MockTransport(handler)


def test_build_embedder_without_cpu_fallback(tmp_path: Path) -> None:
    settings = make_settings(
        tmp_path, embeddings_url="http://gpu.test", embeddings_cpu_fallback=False
    )
    embedder = build_embedder(settings)
    assert embedder._local_factory is None, "модель на CPU не создаётся вовсе"
    assert embedder.provider == "remote"
    # Без сервиса настройка не действует: CPU сервера — единственный провайдер.
    plain = build_embedder(make_settings(tmp_path, embeddings_cpu_fallback=False))
    assert plain._local_factory is not None
    assert plain.provider == "local"


def test_rejected_token_is_named() -> None:
    remote = HttpEmbedder("http://gpu.test", "fake-berta", transport=token_rejected())
    assert remote.health()["device"] == "cuda"
    with pytest.raises(EmbeddingServiceError, match="EMBEDDINGS_TOKEN"):
        remote.encode(["текст"], "topic")


def test_no_cpu_fallback_and_back() -> None:
    now = [0.0]
    broken = HttpEmbedder(
        "http://gpu.test", "fake-berta", token="test-token", transport=gpu_node(fail=True)
    )
    embedder = CachingEmbedder(
        "fake-berta", local=None, remote=broken, retry_after=60, clock=lambda: now[0]
    )
    with pytest.raises(EmbeddingsUnavailableError):
        embedder.encode(["Первый текст про ДТП."], "topic")
    assert embedder.provider == "off"
    assert not embedder.available
    note = embedder.notes()[0]
    assert note.level == "warning"
    assert "Модель смысла недоступна" in note.message
    assert "по словам и признакам события" in note.message
    with pytest.raises(EmbeddingsUnavailableError):  # пока сервис «лежит» — сразу «недоступно»
        embedder.encode(["Второй текст."], "topic")
    # Через минуту — снова пробуем сервис; он ожил.
    embedder._remote = HttpEmbedder(
        "http://gpu.test", "fake-berta", token="test-token", transport=gpu_node()
    )
    now[0] = 61.0
    assert embedder.provider == "remote"  # снова «доступна» — до первой ошибки
    assert embedder.encode(["Третий текст про гололёд."], "topic").shape == (1, 256)
    assert embedder.provider == "remote"
    assert embedder.notes()[0].level == "info"


async def test_warmup_without_embeddings(tmp_path: Path) -> None:
    local = FakeEmbedder()
    remote = HttpEmbedder("http://gpu.test", "fake-berta", transport=token_rejected())
    matcher = make_funnel(
        make_settings(tmp_path), embedder=local, remote=remote, cpu_fallback=False
    )
    await matcher.warmup()  # не падает и не грузит модель на CPU
    assert not local.calls
    assert not matcher.embedder.warm
    assert matcher.embedder.provider == "off"
    assert matcher.warmup_seconds is not None


async def test_search_without_embeddings_is_lexical(tmp_path: Path) -> None:
    """Сервис эмбеддингов не принял токен, отката на CPU нет: поиск идёт по словам и
    признакам события, модель на CPU не грузится, на странице стенда — предупреждение."""
    local = FakeEmbedder()

    def factory(settings: Settings) -> FunnelMatcher:
        remote = HttpEmbedder("http://gpu.test", "fake-berta", transport=token_rejected())
        return make_funnel(settings, embedder=local, remote=remote, cpu_fallback=False)

    settings = make_settings(
        tmp_path, embeddings_url="http://gpu.test", embeddings_cpu_fallback=False
    )
    harness = make_app(settings, matcher_factory=factory)
    async with open_client(harness.app) as client:
        headers = await new_workspace(client)
        search_id, events = await run_search(client, headers)
        assert events[-1]["event"] == "finished"
        selection = (await client.get(f"/api/searches/{search_id}", headers=headers)).json()
        assert_contract("Selection", selection)
        stand = (await client.get("/api/stand")).json()
    assert not local.calls, "модель на CPU не загружалась"
    assert selection["counts"]["total"] > 0
    for item in selection["items"]:
        check_item(item)
        if item["method"] == "ai-unchecked":  # нейросети в тесте нет
            assert item["relevance"] == "disputed"
    assert stand["overall"] == "degraded"
    first = stand["events"][0]
    assert first["level"] == "warning"
    assert "Модель смысла недоступна" in first["message"]
    assert "EMBEDDINGS_TOKEN" in first["message"]


class CardCountingLlm(FakeLlm):
    """FakeLlm, который считает отдельно запросы карточки темы."""

    def __init__(self) -> None:
        super().__init__()
        self.card_calls = 0

    def chat_json(self, system: str, user: str, schema: dict[str, Any]) -> str:
        if "verdicts" not in schema.get("properties", {}):
            self.card_calls += 1
        return super().chat_json(system, user, schema)


async def test_ai_topic_card_without_embeddings(tmp_path: Path) -> None:
    """Нейросеть есть, сервис эмбеддингов не принял токен, отката на CPU нет: карточка темы
    от нейросети собирается без смысловой сверки «не то», поиск не падает, нейросеть
    о карточке спрашивается один раз (стенд 20.10: поиск по теме прерывался ошибкой)."""
    clear_card_cache()
    local = FakeEmbedder()
    llm = CardCountingLlm()

    def factory(settings: Settings) -> FunnelMatcher:
        remote = HttpEmbedder("http://gpu.test", "fake-berta", transport=token_rejected())
        return make_funnel(settings, embedder=local, remote=remote, llm=llm, cpu_fallback=False)

    settings = make_settings(
        tmp_path, embeddings_url="http://gpu.test", embeddings_cpu_fallback=False
    )
    harness = make_app(settings, matcher_factory=factory, judge=AvailableJudge())
    async with open_client(harness.app) as client:
        headers = await new_workspace(client)
        search_id, events = await run_search(client, headers)
        assert events[-1]["event"] == "finished"
        topic = next(e["data"]["topic"] for e in events if e["event"] == "topic")
        selection = (await client.get(f"/api/searches/{search_id}", headers=headers)).json()
    assert topic["madeBy"] == "ai"
    assert "пожарная безопасность" in [f["label"].lower() for f in topic["excludes"]]
    assert llm.card_calls == 1
    assert not local.calls, "модель на CPU не загружалась"
    assert selection["counts"]["total"] > 0
