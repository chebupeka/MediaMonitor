"""Источники пользователя (ТЗ-Ф15, Ф16; П9; НТ-7; НТ-14; НТ-17) и импорт из Excel/CSV."""

from __future__ import annotations

import io
import re
from datetime import UTC, datetime
from pathlib import Path

import httpx
from openpyxl import Workbook, load_workbook

from mediamonitor.adapters.base import RawPost, SourceInfo
from mediamonitor.adapters.fixture import FixtureStore, Recording
from mediamonitor.normalize.urls import PostRef
from tests.conftest import (
    Harness,
    make_app,
    make_settings,
    new_workspace,
    open_client,
    run_search,
)


async def test_first_run_and_catalog(client: httpx.AsyncClient) -> None:
    headers = await new_workspace(client)
    start = (await client.get("/api/start", headers=headers)).json()
    assert start["firstRun"] is True
    assert start["sourcesCount"] == 0
    assert start["defaultMode"] == "everywhere"
    assert start["defaultPeriod"]["preset"] == "7d"
    # Лимиты сервер проверяет сам; в контракте ветки интерфейса их в StartState нет.
    assert "limits" not in start
    catalog = (await client.get("/api/catalog", headers=headers)).json()
    assert catalog
    assert all(c["regional"] for c in catalog)
    assert not any("закрыт" in c["description"] for c in catalog)
    added = (
        await client.post(
            "/api/sources/from-catalog",
            json={"ids": [c["id"] for c in catalog[:3]] * 2},
            headers=headers,
        )
    ).json()
    assert len(added) == 3
    assert all(s["origin"] == "catalog" for s in added)
    start = (await client.get("/api/start", headers=headers)).json()
    assert start["firstRun"] is False
    assert start["sourcesCount"] == 3


async def test_check_states(client: httpx.AsyncClient) -> None:
    headers = await new_workspace(client)

    async def state(url: str) -> str:
        return str(
            (await client.post("/api/sources/check", json={"url": url}, headers=headers)).json()[
                "state"
            ]
        )

    assert await state("https://vk.com/demo_minusinsk") == "ok"
    assert await state("https://vk.com/demo_kodinsk") == "closed"
    assert await state("https://vk.com/demo_news_sibir") == "not-regional"
    assert await state("https://example.ru/news") == "not-supported"
    assert await state("https://vk.com/nobody_here") == "not-found"
    post = (
        await client.post(
            "/api/sources/check",
            json={"url": "https://max.ru/demo_achinsk24/1520"},
            headers=headers,
        )
    ).json()
    assert post["state"] == "post-link"
    assert post["preview"]["url"] == "https://max.ru/demo_achinsk24"
    await client.post(
        "/api/sources", json={"url": "https://vk.com/demo_minusinsk"}, headers=headers
    )
    assert await state("https://m.vk.com/demo_minusinsk") == "duplicate"


async def test_add_update_remove_restore(client: httpx.AsyncClient) -> None:
    headers = await new_workspace(client)
    closed = await client.post(
        "/api/sources", json={"url": "https://vk.com/demo_kodinsk"}, headers=headers
    )
    assert closed.status_code == 422
    assert closed.json()["code"] == "invalid"
    source = (
        await client.post(
            "/api/sources",
            json={
                "url": "https://vk.com/demo_minusinsk",
                "name": " Минусинск ",
                "type": "municipal",
                "district": None,
            },
            headers=headers,
        )
    ).json()
    assert source["name"] == "Минусинск"
    assert source["type"] == "municipal"
    assert source["district"] is None
    assert source["inMySources"] is True
    assert source["check"]["state"] == "ok"
    updated = (
        await client.patch(
            f"/api/sources/{source['id']}", json={"district": "Минусинск"}, headers=headers
        )
    ).json()
    assert updated["district"] == "Минусинск"
    assert updated["name"] == "Минусинск"
    assert (await client.delete(f"/api/sources/{source['id']}", headers=headers)).status_code == 204
    assert (await client.get("/api/sources", headers=headers)).json() == []
    restored = await client.post(f"/api/sources/{source['id']}/restore", headers=headers)
    assert restored.status_code == 200
    assert restored.json()["district"] == "Минусинск"
    again = await client.post(f"/api/sources/{source['id']}/restore", headers=headers)
    assert again.status_code == 404


async def test_closed_group_in_list_is_partial(harness: Harness, client: httpx.AsyncClient) -> None:
    """Сц. ТЗ 5 без имитации: группа из списка закрылась → VK «частично» (НТ-17)."""
    from mediamonitor.adapters.fixture import FixtureStore
    from mediamonitor.services import sources as source_service
    from mediamonitor.services.workspace import find_workspace

    headers = await new_workspace(client)
    catalog = (await client.get("/api/catalog", headers=headers)).json()
    vk = [c["id"] for c in catalog if c["platform"] == "vk"][:5]
    await client.post("/api/sources/from-catalog", json={"ids": vk}, headers=headers)
    # Группа была открыта, когда её добавили, и закрылась потом.
    closed = FixtureStore.load().sources["vk-boguchany"]
    async with harness.app.state.ctx.sessions.begin() as session:
        workspace = await find_workspace(session, headers["X-Workspace-Key"])
        assert workspace is not None
        directory = await source_service.upsert_source(session, closed)
        await source_service.attach(
            session, workspace.id, directory, "manual", source_service.ok_check()
        )
    report = (await client.post("/api/sources/recheck", headers=headers)).json()
    assert report["checked"] == 6
    assert report["ok"] == 5
    assert [p["state"] for p in report["problems"]] == ["closed"]
    _, events = await run_search(client, headers, mode="my-sources")
    selection = events[-1]["data"]["selection"]
    statuses = {s["platform"]: s for s in selection["statuses"]}
    assert selection["conditions"]["sourceCount"] == 6
    assert statuses["vk"]["state"] == "partial"
    assert statuses["vk"]["problem"]["code"] == "sources-closed"
    assert "Новости Богучан" in statuses["vk"]["problem"]["message"]
    assert statuses["ok"]["state"] == "nothing"
    assert statuses["ok"]["progress"] == {
        "done": 0,
        "total": 0,
        "unit": "sources",
    }  # «нет источников в списке»


async def test_import_xlsx(client: httpx.AsyncClient) -> None:
    headers = await new_workspace(client)
    await client.post("/api/sources", json={"url": "https://vk.com/demo_tv_kraya"}, headers=headers)
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    rows = [
        ("Название", "Ссылка на группу", "Комментарий"),
        ("Лесосибирск", "  https://vk.com/demo_lesosib  ", "пробелы"),
        ("Минусинск", "https://m.vk.com/demo_minusinsk", "мобильная"),
        ("Сайт", "https://example.ru/news", ""),
        ("Телеканал", "https://vk.com/demo_tv_kraya", "уже есть"),
        (None, None, None),
        ("Кодинск", "https://vk.com/demo_kodinsk", "закрыта"),
        ("Пост", "https://vk.com/wall-100104_60210", ""),
        ("Повтор", "https://vk.com/demo_lesosib", "второй раз"),
    ]
    for row in rows:
        sheet.append(row)
    buffer = io.BytesIO()
    workbook.save(buffer)
    preview = (
        await client.post(
            "/api/sources/import/preview",
            files={"file": ("источники.xlsx", buffer.getvalue(), "application/octet-stream")},
            headers=headers,
        )
    ).json()
    assert [c["role"] for c in preview["columns"]] == ["name", "url", "ignored"]
    assert preview["headerRow"] == 1
    assert preview["totalRows"] == 9
    outcomes = [(r["row"], r["outcome"]["kind"]) for r in preview["rows"]]
    assert outcomes == [
        (2, "add"),
        (3, "add"),
        (4, "skip"),
        (5, "duplicate"),
        (6, "empty"),
        (7, "skip"),
        (8, "skip"),
        (9, "duplicate"),
    ]
    notes = {
        r["row"]: r["outcome"].get("note") for r in preview["rows"] if r["outcome"]["kind"] == "add"
    }
    assert notes == {2: "убрали лишние пробелы", 3: "мобильная ссылка — привели к обычной"}
    assert preview["counts"] == {"add": 2, "duplicate": 2, "skip": 3, "empty": 1}
    # Разбор ничего не добавил.
    assert len((await client.get("/api/sources", headers=headers)).json()) == 1
    report = (
        await client.post(f"/api/sources/import/{preview['previewId']}", headers=headers)
    ).json()
    assert len(report["added"]) == 2
    assert all(s["origin"] == "import" for s in report["added"])
    assert len(report["duplicates"]) == 2
    assert len(report["skipped"]) == 3
    assert len((await client.get("/api/sources", headers=headers)).json()) == 3
    gone = await client.post(f"/api/sources/import/{preview['previewId']}", headers=headers)
    assert gone.status_code == 404


async def test_import_rejects_other_files(client: httpx.AsyncClient) -> None:
    headers = await new_workspace(client)
    for name in ("фото.png", "старый.xls"):
        response = await client.post(
            "/api/sources/import/preview",
            files={"file": (name, b"data", "application/octet-stream")},
            headers=headers,
        )
        assert response.status_code == 422
        assert name in response.json()["message"]
    broken = await client.post(
        "/api/sources/import/preview",
        files={"file": ("битый.xlsx", b"not a zip", "application/octet-stream")},
        headers=headers,
    )
    assert broken.status_code == 422


# ---------------------------------------------------------------- бэклог п. 47–48


async def test_default_mode_is_remembered(client: httpx.AsyncClient) -> None:
    """Первый запуск — «Везде», дальше — последний выбор человека (п. 47)."""
    headers = await new_workspace(client)

    async def default_mode() -> str:
        return str((await client.get("/api/start", headers=headers)).json()["defaultMode"])

    assert await default_mode() == "everywhere"
    catalog = (await client.get("/api/catalog", headers=headers)).json()
    await client.post(
        "/api/sources/from-catalog", json={"ids": [catalog[0]["id"]]}, headers=headers
    )
    await run_search(client, headers, mode="my-sources")
    assert await default_mode() == "my-sources"
    await run_search(client, headers, mode="everywhere")
    assert await default_mode() == "everywhere"
    await run_search(client, headers, mode="my-sources")
    # Список опустел — искать «по моим» нечему, снова «Везде» (НТ-14).
    await client.delete(f"/api/sources/{catalog[0]['id']}", headers=headers)
    assert await default_mode() == "everywhere"
    other = await new_workspace(client)  # у другого рабочего места — свой выбор
    assert (await client.get("/api/start", headers=other)).json()["defaultMode"] == "everywhere"


FEDERAL = SourceInfo(
    platform="vk",
    owner_id="-900001",
    slug="fed_auto_news",
    name="Автоновости страны",
    url="https://vk.com/fed_auto_news",
    subscribers=500_000,
    type="media",
    district=None,
    regional=False,
)


def federal_store() -> FixtureStore:
    texts = [
        "С 1 марта меняются правила техосмотра. Владельцам машин стоит записаться заранее.",
        "В Канске открыли новый пункт техосмотра. Очередей пока нет.",
    ]
    posts = [
        RawPost(
            ref=PostRef("vk", "-900001", str(10 + i)),
            source=FEDERAL,
            text=text,
            title=None,
            published_at=datetime(2026, 10, 12, 5, i, tzinfo=UTC),
            time_reliability="exact",
            edited=False,
            repost_of=None,
            has_media=False,
            views=1000,
            reactions=10,
        )
        for i, text in enumerate(texts)
    ]
    return FixtureStore(
        sources={"fed": FEDERAL},
        recordings=[
            Recording(
                id="techosmotr",
                regex=[re.compile("техосмотр")],
                shingle_index=set(),
                min_shingle_hits=0,
                meaning_id=None,
                posts=posts,
            )
        ],
        posts={p.ref.id: p for p in posts},
        everywhere_pool={"vk": 2, "ok": 1, "max": 1},
    )


async def test_my_sources_are_not_filtered_by_region(tmp_path: Path) -> None:
    """Р4, п. 48: «Мои источники» — регион не проверяем, причина «в вашем списке»."""
    harness = make_app(make_settings(tmp_path), store=federal_store())
    async with open_client(harness.app) as client:
        headers = await new_workspace(client)
        check = (
            await client.post("/api/sources/check", json={"url": FEDERAL.url}, headers=headers)
        ).json()
        assert check["state"] == "not-regional"
        await client.post("/api/sources", json={"url": FEDERAL.url}, headers=headers)

        search_id, events = await run_search(client, headers, "техосмотр", mode="my-sources")
        items = {i["id"]: i for i in events[-1]["data"]["selection"]["items"]}
        assert set(items) == {"vk:-900001_10", "vk:-900001_11"}
        assert items["vk:-900001_10"]["regionReason"] == {
            "regionalSource": False,
            "inMySources": True,
            "places": [],
        }
        assert items["vk:-900001_11"]["regionReason"]["places"] == ["Канск"]
        assert items["vk:-900001_11"]["regionReason"]["inMySources"] is True
        response = await client.get(f"/api/searches/{search_id}/export.xlsx", headers=headers)
        sheet = load_workbook(io.BytesIO(response.content))["Публикации"]
        reasons = {sheet.cell(row=r, column=11).value for r in range(2, sheet.max_row + 1)}
        assert reasons == {"в вашем списке", "в вашем списке · упоминается: Канск"}

        # «Везде»: прежнее правило — краевой источник или место края в тексте.
        _, events = await run_search(client, headers, "техосмотр", mode="everywhere")
        items = {i["id"]: i for i in events[-1]["data"]["selection"]["items"]}
        assert set(items) == {"vk:-900001_11"}
        assert items["vk:-900001_11"]["regionReason"] == {
            "regionalSource": False,
            "inMySources": False,
            "places": ["Канск"],
        }


async def test_list_only_posts_from_fixtures(client: httpx.AsyncClient) -> None:
    """Записи listOnly мокапа: посты некраевых источников приходят только по списку (п. 48)."""
    headers = await new_workspace(client)
    for url in ("https://vk.com/demo_news_sibir", "https://vk.com/demo_avto_rf"):
        added = await client.post("/api/sources", json={"url": url}, headers=headers)
        assert added.status_code == 201, added.text
    _, events = await run_search(client, headers, mode="my-sources")
    items = {i["id"]: i for i in events[-1]["data"]["selection"]["items"]}
    for key in ("vk:-100201_57102", "vk:-100202_310577"):
        assert key in items
        assert items[key]["regionReason"] == {
            "regionalSource": False,
            "inMySources": True,
            "places": [],
        }
    _, events = await run_search(client, headers, mode="everywhere")
    ids = {i["id"] for i in events[-1]["data"]["selection"]["items"]}
    assert "vk:-100201_57102" not in ids
    assert "vk:-100202_310577" not in ids
