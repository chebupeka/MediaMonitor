"""Выгрузка XLSX (ТЗ-Ф37, Ф38; НТ-10, НТ-11, НТ-12, НТ-20).

Правила (решения из матрицы, п. 1):
- в файл попадает вся подборка минус исключённые; фильтр вида на файл не влияет, пока
  не выбрано явно «только эти метки» или «только без даты»;
- группы выгружаются строкой на каждое размещение, с колонкой «Группа»;
- пост без даты — пустая дата и пометка «дата неизвестна» в соседней колонке;
- даты — настоящие даты Excel по красноярскому времени, подпись пояса — в заголовке;
- ячейка, начинающаяся с = + - @ (а также табуляции и перевода строки), — только текст;
- колонка «Ссылка» — гиперссылка, построенная сервисом из ключа поста.

Предпросмотр и файл собираются из одних и тех же строк — «файл совпадает с экраном».
"""

from __future__ import annotations

import io
import re
from collections.abc import Sequence
from datetime import datetime

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE, Cell
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from mediamonitor.contract import (
    RELEVANCE_ORDER,
    ExportCell,
    ExportColumn,
    ExportOptions,
    ExportPreview,
    ExportSheetTable,
    ExportSummarySheet,
    ResultItem,
    Selection,
    SummaryEntry,
)
from mediamonitor.errors import required
from mediamonitor.labels import (
    METHOD_LABEL,
    MODE_LABEL,
    PLACEMENT_TYPE_LABEL,
    PLATFORM_NAME,
    QUERY_KIND_LABEL,
    RELEVANCE_LABEL,
    format_date_time,
    format_period,
    platform_status_text,
    region_reason_text,
    truncate_text,
)
from mediamonitor.normalize.urls import is_safe_post_url
from mediamonitor.services.selection import compute_counts
from mediamonitor.timeutil import KRSK, TIME_ZONE_LABEL_FULL, excel_local, parse_iso

PREVIEW_ROWS = 50
XLSX_CELL_LIMIT = 32_767
DANGEROUS_START = ("=", "+", "-", "@", "\t", "\r", "\n")
COLUMNS: list[ExportColumn] = [
    ExportColumn(id="n", title="№"),
    ExportColumn(id="platform", title="Площадка"),
    ExportColumn(id="source", title="Источник"),
    ExportColumn(id="headline", title="Заголовок или начало поста"),
    ExportColumn(id="date", title="Дата публикации (время красноярское)"),
    ExportColumn(id="dateNote", title="Пометка о дате"),
    ExportColumn(id="url", title="Ссылка"),
    ExportColumn(id="excerpt", title="Выдержка"),
    ExportColumn(id="relevance", title="Метка"),
    ExportColumn(id="method", title="Как отобрано"),
    ExportColumn(id="region", title="Почему краевой"),
    ExportColumn(id="placement", title="Тип размещения"),
    ExportColumn(id="group", title="Группа"),
    ExportColumn(id="fetched", title="Получено"),
]
WIDTHS = [6, 14, 28, 50, 20, 30, 42, 70, 18, 20, 30, 16, 36, 18]
UNDATED_NOTE = "дата неизвестна: в период не засчитана"
APPROXIMATE_NOTE = "≈ время приблизительное: вычислено по номеру поста MAX"


def safe_text(value: str) -> str:
    """Для предпросмотра: значение, начинающееся с = + - @, показываем с апострофом (НТ-10)."""
    return f"'{value}" if value.startswith(DANGEROUS_START) else value


def export_items(selection: Selection, options: ExportOptions | None) -> list[ResultItem]:
    relevance = set(options.only_relevance) if options and options.only_relevance else None
    only_undated = bool(options and options.only_undated)
    items = [
        i
        for i in selection.items
        if not i.excluded
        and (relevance is None or i.relevance in relevance)
        and (not only_undated or i.publication.published_at is None)
    ]
    # Новые сверху, без даты — в конце (как на экране).
    items.sort(key=lambda i: i.publication.published_at or "", reverse=True)
    return items


def file_name(selection: Selection) -> str:
    conditions = selection.conditions
    day = (conditions.finished_at or conditions.started_at)[:10]
    label = re.sub(r"[^\w]+", "_", conditions.label, flags=re.UNICODE).strip("_")[:40] or "подборка"
    return f"Медиамонитор_{label}_{day}.xlsx"


def _date_note(item: ResultItem) -> str:
    publication = item.publication
    if publication.published_at is None:
        return UNDATED_NOTE
    if publication.time_reliability == "approximate":
        return APPROXIMATE_NOTE
    return ""


def _group_titles(selection: Selection) -> dict[str, str]:
    return {g.id: g.title for g in selection.groups}


def _rule(options: ExportOptions | None) -> str:
    if options and options.only_undated:
        return (
            "В файл попадут только записи без даты — то, что видно на экране с фильтром «Дата "
            "неизвестна». Исключённые вами записи не попадут."
        )
    if options and options.only_relevance:
        labels = ", ".join(RELEVANCE_LABEL[r] for r in options.only_relevance)
        return (
            f"В файл попадут только записи с метками: {labels}. Исключённые вами записи не попадут."
        )
    return (
        "В файл попадает вся подборка, кроме исключённых вами записей. Фильтр на экране на файл не "
        "влияет. Группы выгружаются построчно: у каждого размещения своя строка и колонка «Группа»."
    )


def summary_sheet(
    selection: Selection, rows: Sequence[ResultItem], data_note: str
) -> ExportSummarySheet:
    c = selection.conditions
    counts = compute_counts(rows)
    period = format_period(
        datetime.fromisoformat(c.period.from_).date(), datetime.fromisoformat(c.period.to).date()
    )
    sources = (
        f"{MODE_LABEL[c.mode]}: {c.source_count or 0}"
        if c.mode == "my-sources"
        else MODE_LABEL[c.mode]
    )
    entries = [
        SummaryEntry(label="Запрос", value=truncate_text(c.text, 300)),
        SummaryEntry(
            label="Что искали",
            value=QUERY_KIND_LABEL[c.kind] + (f" · {c.meaning_label}" if c.meaning_label else ""),
        ),
        SummaryEntry(label="Период", value=f"{period} ({TIME_ZONE_LABEL_FULL})"),
        SummaryEntry(label="Регион", value=c.region.name),
        SummaryEntry(label="Источники", value=sources),
        SummaryEntry(
            label="Площадки",
            value="; ".join(
                f"{PLATFORM_NAME[s.platform]}: {platform_status_text(s)}"
                for s in selection.statuses
            ),
        ),
        SummaryEntry(
            label="Поиск выполнен", value=format_date_time(parse_iso(c.finished_at or c.started_at))
        ),
        SummaryEntry(label="Записей в файле", value=str(len(rows))),
        SummaryEntry(label="Из них без даты", value=str(counts.undated)),
        SummaryEntry(label="Исключено вами", value=str(selection.counts.excluded)),
        SummaryEntry(
            label="По меткам",
            value=" · ".join(
                f"{RELEVANCE_LABEL[r]}: {getattr(counts.by_relevance, r)}" for r in RELEVANCE_ORDER
            ),
        ),
        SummaryEntry(label="Время", value=TIME_ZONE_LABEL_FULL),
        SummaryEntry(label="Данные", value=data_note),
    ]
    return ExportSummarySheet(name="Сводка", entries=entries)


def _row(
    item: ResultItem, number: int, groups: dict[str, str]
) -> list[str | int | datetime | None]:
    """Значения строки листа «Публикации»: даты — datetime, остальное — текст или число."""
    p = item.publication
    return [
        number,
        PLATFORM_NAME[p.platform],
        p.source.name,
        p.headline,
        excel_local(parse_iso(p.published_at)) if p.published_at else None,
        _date_note(item),
        p.url,
        item.excerpt.text if item.excerpt else "текст отсутствует",
        RELEVANCE_LABEL[item.relevance],
        METHOD_LABEL[item.method],
        region_reason_text(item.region_reason),
        PLACEMENT_TYPE_LABEL[item.placement_type] if item.placement_type else "",
        groups.get(item.group_id, "") if item.group_id else "",
        excel_local(parse_iso(p.fetched_at)),
    ]


def _preview_cell(value: str | int | datetime | None) -> ExportCell:
    if isinstance(value, datetime):
        return format_date_time(value.replace(tzinfo=KRSK))
    if isinstance(value, str):
        return safe_text(value)
    return value


def build_preview(
    selection: Selection, options: ExportOptions | None, data_note: str
) -> ExportPreview:
    rows = export_items(selection, options)
    groups = _group_titles(selection)
    table = [
        [_preview_cell(v) for v in _row(item, n + 1, groups)]
        for n, item in enumerate(rows[:PREVIEW_ROWS])
    ]
    applied = list(options.only_relevance) if options and options.only_relevance else None
    return ExportPreview(
        search_id=selection.search_id,
        file_name=file_name(selection),
        row_count=len(rows),
        excluded_count=selection.counts.excluded,
        rule=_rule(options),
        applied_filter=applied,
        only_undated=bool(options and options.only_undated),
        summary=summary_sheet(selection, rows, data_note),
        publications=ExportSheetTable(
            name="Публикации", columns=COLUMNS, rows=table, total_rows=len(rows)
        ),
        time_zone_note=TIME_ZONE_LABEL_FULL,
    )


# ---------------------------------------------------------------- файл


def _text_cell(cell: Cell, value: str) -> None:
    """Записать строку только как текст: формулой она не станет ни в одной программе (НТ-10)."""
    clean = ILLEGAL_CHARACTERS_RE.sub("", value)[:XLSX_CELL_LIMIT]
    cell.value = clean
    cell.data_type = "s"
    if clean.startswith(DANGEROUS_START):
        cell.quotePrefix = True


def _write(sheet: Worksheet, row: int, column: int, value: str | int | datetime | None) -> Cell:
    cell = sheet.cell(row=row, column=column)
    if not isinstance(cell, Cell):  # объединённых ячеек в выгрузке нет
        raise TypeError("Ожидалась обычная ячейка")
    if isinstance(value, str):
        _text_cell(cell, value)
    elif isinstance(value, datetime):
        cell.value = value
        cell.number_format = "DD.MM.YYYY HH:MM"
    else:
        cell.value = value
    return cell


def build_xlsx(selection: Selection, options: ExportOptions | None, data_note: str) -> bytes:
    rows = export_items(selection, options)
    groups = _group_titles(selection)
    workbook = Workbook()
    sheet = workbook.active
    sheet = required(sheet, "sheet")
    sheet.title = "Публикации"
    bold = Font(bold=True)
    for index, column in enumerate(COLUMNS, start=1):
        cell = _write(sheet, 1, index, column.title)
        cell.font = bold
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        sheet.column_dimensions[cell.column_letter].width = WIDTHS[index - 1]
    url_column = next(i for i, c in enumerate(COLUMNS, start=1) if c.id == "url")
    for offset, item in enumerate(rows, start=2):
        for index, value in enumerate(_row(item, offset - 1, groups), start=1):
            cell = _write(sheet, offset, index, value)
            if index == url_column and isinstance(value, str) and is_safe_post_url(value):
                cell.hyperlink = value
                cell.style = "Hyperlink"
                cell.data_type = "s"
    sheet.freeze_panes = "A2"
    last_column = get_column_letter(len(COLUMNS))
    sheet.auto_filter.ref = f"A1:{last_column}{max(1, len(rows) + 1)}"

    summary = workbook.create_sheet("Сводка")
    sheet_data = summary_sheet(selection, rows, data_note)
    for offset, entry in enumerate(sheet_data.entries, start=1):
        _write(summary, offset, 1, entry.label).font = bold
        _write(summary, offset, 2, entry.value).alignment = Alignment(
            wrap_text=True, vertical="top"
        )
    summary.column_dimensions["A"].width = 22
    summary.column_dimensions["B"].width = 100

    workbook.properties.creator = "Медиамонитор"
    workbook.properties.title = file_name(selection)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
