"""Сквозная проверка живого сервера: /health → поиск → поток SSE → подборка → предпросмотр → XLSX.

Запуск: uv run python scripts/smoke.py [адрес]   (по умолчанию http://127.0.0.1:8000)
Работает с данными DATA_SOURCE=fixtures (записанные примеры).
"""

from __future__ import annotations

import io
import json
import sys
import time
from typing import Any

import httpx
from openpyxl import load_workbook

PERIOD = {"preset": "custom", "from": "2026-09-20", "to": "2026-10-19"}


def main(base: str) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")  # консоль Windows — cp1251
    with httpx.Client(base_url=base, timeout=60) as client:
        health = client.get("/health").json()
        print("health:", health)
        assert health["status"] == "ok", health

        key = client.post("/api/workspaces").json()["key"]
        headers = {"X-Workspace-Key": key}
        body = {"text": "Безопасность дорожного движения", "period": PERIOD, "mode": "everywhere"}
        started = time.perf_counter()
        ticket = client.post("/api/searches", json=body, headers=headers)
        assert ticket.status_code == 202, ticket.text
        search_id = ticket.json()["searchId"]

        events: list[dict[str, Any]] = []
        first_items: float | None = None
        with client.stream("GET", f"/api/searches/{search_id}/events", headers=headers) as stream:
            assert stream.headers["content-type"].startswith("text/event-stream")
            for line in stream.iter_lines():
                if line.startswith("data: "):
                    event = json.loads(line[6:])
                    events.append(event)
                    if event["type"] == "items" and first_items is None:
                        first_items = time.perf_counter() - started
        took = time.perf_counter() - started
        types = [e["type"] for e in events]
        first = first_items or 0.0
        print(f"поток: {len(events)} событий за {took:.1f} с; первые записи через {first:.1f} с")
        print("  порядок:", " → ".join(dict.fromkeys(types)))
        assert types[0] == "started" and types[-1] == "finished", types

        selection = client.get(f"/api/searches/{search_id}", headers=headers).json()
        statuses = " · ".join(
            f"{s['platform']}: {s['state']} {s['found']}" for s in selection["statuses"]
        )
        print(f"подборка: {selection['counts']['total']} записей; {statuses}")

        preview = client.get(f"/api/searches/{search_id}/export/preview", headers=headers).json()
        print(f"предпросмотр: «{preview['fileName']}», строк {preview['rowCount']}")
        assert preview["rowCount"] == selection["counts"]["total"]

        response = client.get(f"/api/searches/{search_id}/export.xlsx", headers=headers)
        assert response.status_code == 200, response.text
        workbook = load_workbook(io.BytesIO(response.content))
        rows = workbook["Публикации"].max_row - 1
        print(f"XLSX: {len(response.content)} байт, листы {workbook.sheetnames}, строк {rows}")
        assert rows == preview["rowCount"]
    print("ГОТОВО: сквозной сценарий пройден")
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv[1] if len(sys.argv) > 1 else "http://127.0.0.1:8000"))
